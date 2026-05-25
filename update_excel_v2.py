import pandas as pd
import numpy as np

# Load the updated NDVI data (pixel level)
df_ndvi = pd.read_excel('Plan de pastoreo_updated.xlsx', sheet_name='NDVI')

# Calculate area for each section
areas = df_ndvi.groupby('id')['area'].sum()
print("Areas (ha):", areas)

# The formulas in other sheets usually depend on the NDVI average per section
# I will create a summary sheet for the user and check if I can update the existing sheets.

# However, the instruction says "Complete the Excel workbook".
# Given the complexity of Excel formulas, the best is to keep the "NDVI" sheet
# and potentially a "Resumen_GIS" sheet.
# But I should try to make it work with the existing structure.
# The existing structure had C1..C14. I now have C1..C3.

# Let's try to update "kg MS" sheet to only show C1, C2, C3.
orig_xlsx = pd.ExcelFile('Plan de pastoreo.xlsx')
sheets = {name: orig_xlsx.parse(name) for name in orig_xlsx.sheet_names}

# We can't easily update all formulas because they are complex.
# But we can provide the data in the way the workbook expects.
# If I change the NDVI sheet to have 3 IDs (1, 2, 3),
# and the other sheets reference IDs, they might break if they expect 1 to 14.

# Let's see what IDs are used in the original NDVI sheet.
df_orig_ndvi = pd.read_excel('Plan de pastoreo.xlsx', sheet_name='NDVI')
print("Original NDVI unique IDs:", df_orig_ndvi['id'].unique())
# It had 1..14.

# If I map my 3 sections to IDs 1, 2, 3, then columns C1, C2, C3 in "kg MS" should work
# if they are driven by the NDVI sheet.
# But usually they are Pivot Tables or direct references.

# Let's just ensure the NDVI sheet is perfect.
# I will use openpyxl to keep existing formulas in other sheets.

import openpyxl

wb = openpyxl.load_workbook('Plan de pastoreo.xlsx')
ws_ndvi = wb['NDVI']

# Clear existing data in NDVI sheet
for row in ws_ndvi.iter_rows(min_row=2):
    for cell in row:
        cell.value = None

# Write new data
for i, row in df_ndvi.iterrows():
    for j, val in enumerate(row):
        ws_ndvi.cell(row=i+2, column=j+1, value=val)

wb.save('Plan de pastoreo_updated_v2.xlsx')
print("Saved v2 with openpyxl preserving other sheets.")
