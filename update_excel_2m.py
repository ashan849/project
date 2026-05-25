import pandas as pd
import numpy as np
import rasterio
import geopandas as gpd
import glob
import openpyxl

# Load subdivisions and masks
subs = gpd.read_file("subdivisions.gpkg")

def extract_pixel_data_2m():
    rows = []

    # Load all 2m NDVI rasters
    ndvi_files = sorted(glob.glob("rasters_2m/ndvi_*.tif"))
    ndvi_data = {}
    for f in ndvi_files:
        date_str = f.split("_")[1].split(".")[0]
        with rasterio.open(f) as src:
            ndvi_data[date_str] = src.read(1)

    with rasterio.open("rasters_2m/slope_25830.tif") as src:
        slope_data = src.read(1)
    with rasterio.open("rasters_2m/distance_to_water.tif") as src:
        dist_data = src.read(1)
    with rasterio.open("rasters_2m/grass_mask.tif") as src:
        grass_data = src.read(1)
        affine = src.transform

    for r in range(grass_data.shape[0]):
        for c in range(grass_data.shape[1]):
            # Use a threshold for the resampled mask
            if grass_data[r, c] > 0.5:
                # Get coordinates
                x, y = affine * (c + 0.5, r + 0.5)
                point = gpd.points_from_xy([x], [y])[0]

                # Find which subdivision
                sub_id = np.nan
                for idx, sub in subs.iterrows():
                    if sub.geometry.contains(point):
                        sub_id = sub['id']
                        break

                if np.isnan(sub_id): continue

                pixel_row = {
                    'id': sub_id,
                    'area': 0.0004, # 4 m2 = 0.0004 ha
                    'pendiente': slope_data[r, c],
                    'distancia': dist_data[r, c]
                }

                for date_str, data in ndvi_data.items():
                    pixel_row[f"{date_str} NDVI"] = data[r, c]

                rows.append(pixel_row)

    return pd.DataFrame(rows)

print("Extracting 2m pixel data...")
df_new_ndvi = extract_pixel_data_2m()
print(f"Extracted {len(df_new_ndvi)} pixels.")

# Update the main Excel file
wb = openpyxl.load_workbook('Plan de pastoreo.xlsx')
ws_ndvi = wb['NDVI']

# Clear existing data in NDVI sheet (except header)
for row in ws_ndvi.iter_rows(min_row=2):
    for cell in row:
        cell.value = None

# Write new data
for i, row in df_new_ndvi.iterrows():
    if i % 1000 == 0: print(f"Writing row {i}...")
    for j, val in enumerate(row):
        # Conversion to native types for Excel
        if isinstance(val, (np.float32, np.float64)):
            val = float(val)
        elif isinstance(val, (np.int32, np.int64)):
            val = int(val)
        ws_ndvi.cell(row=i+2, column=j+1, value=val)

wb.save('Plan de pastoreo.xlsx')
print("Excel updated and saved.")
